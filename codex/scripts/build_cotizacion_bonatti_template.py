#!/usr/bin/env python3
"""Turn the real BONATTI quotation workbook into the COTIZACION master template.

The source file is the quotation Arefil sends today ("BONATTI FILTROS
LMR850205-048.xlsx"): a hand-filled document, not a template. This script keeps
its layout, styles, column widths, merges, currency formats and commercial
conditions, and replaces only the data with the placeholders the Report Builder
resolves ({{parameters.*}}, {{rows.*}}, {{summary.*}}).

Everything openpyxl cannot round-trip is reported instead of being invented;
see codex/output/cotizacion-bonatti-e2e.md.

Usage:
    python codex/scripts/build_cotizacion_bonatti_template.py SOURCE.xlsx OUT.xlsx
      [--with-letterhead]

--with-letterhead re-anchors xl/media/image1.jpeg (the Arefil letterhead, a
legacy VML shape openpyxl always drops) as a standard DrawingML picture over
the same 612x795 pt page area. Excel and LibreOffice paint pictures ON TOP of
cell text, so the body of the quotation disappears behind it; the flag exists to
reproduce that experiment, not as the recommended configuration.
"""
from __future__ import annotations

import argparse
import io
import zipfile
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

SHEET = "Hoja1"

# Row 19 of the source is the first quotation line; the renderer repeats any row
# holding {{rows.*}} once per dataset row.
REPEAT_ROW = 19
FIRST_DATA_ROW = 20
LAST_TOTALS_ROW = 35  # rows 20..35 = sample lines + the 3-row totals block

HEADER_CELLS = {
    "E12": "Ciudad Madero, Tamps., a {{parameters.quotation_date}}",
    "B13": "Cliente: {{parameters.customer_name}}",
    "E13": "RFC: {{parameters.customer_rfc}}",
    "B14": "Email: {{parameters.customer_email}}",
    "B15": "Atencion: {{parameters.attention_to}}",
    "C16": "{{parameters.requisition}}",
    # The source column "T/E" (delivery time) has no counterpart in the report
    # contract; the discount the operator captures per product takes its place.
    "E18": "Descuento",
}

REPEAT_CELLS = {
    "A19": "{{rows.row_number}}",
    "B19": "{{rows.quantity}}",
    "C19": "{{rows.sku}}",
    "D19": "{{rows.description}}",
    "E19": "{{rows.discount_amount}}",
    "F19": "{{rows.price}}",
    "G19": "{{rows.subtotal}}",
}

# label, summary placeholder, source row whose style is copied
TOTALS = (
    ("SUBTOTAL", "{{summary.subtotal}}", 33),
    ("DESCUENTO", "{{summary.discount_amount}}", 34),
    ("IVA", "{{summary.tax}}", 34),
    ("TOTAL", "{{summary.total}}", 35),
)

# Merges kept above the repeatable row, plus the RFC line this template adds.
MERGES_ABOVE = ("C1:G10", "E12:G12", "B13:D13", "E13:G13", "B14:D14", "B15:D15", "C16:E16", "C17:G17")
# Everything under the totals block moves up by (16 deleted - 4 inserted) rows.
SHIFT = (LAST_TOTALS_ROW - FIRST_DATA_ROW + 1) - len(TOTALS)
MERGES_BELOW = ("C38:D38", "C39:D39", "C40:D40", "C41:D41", "C42:F42", "C43:D43", "C44:D44",
                "A46:G46", "A47:G47", "A48:G48")
TOTALS_BLOCK = f"B{FIRST_DATA_ROW}:E{FIRST_DATA_ROW + len(TOTALS) - 1}"

LETTERHEAD_WIDTH_PT = 612
LETTERHEAD_HEIGHT_PT = 795


def _shift_range(reference: str, delta: int) -> str:
    start, end = reference.split(":")

    def move(cell: str) -> str:
        column = "".join(character for character in cell if character.isalpha())
        row = int("".join(character for character in cell if character.isdigit()))
        return f"{column}{row + delta}"

    return f"{move(start)}:{move(end)}"


def _copy_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)


def build(source: Path, output: Path, *, with_letterhead: bool) -> None:
    workbook = openpyxl.load_workbook(source)
    sheet: Worksheet = workbook[SHEET]

    heights = {index: dimension.height for index, dimension in sheet.row_dimensions.items()
               if dimension.height is not None}
    totals_styles = [copy(sheet.cell(row, column)) for _, _, row in TOTALS for column in (6, 7)]
    currency_format = sheet["F19"].number_format

    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))

    for coordinate, value in HEADER_CELLS.items():
        cell = sheet[coordinate]
        if coordinate == "E13":  # new line: no style of its own in the source
            _copy_style(sheet["B13"], cell)
        cell.value = value
    for coordinate, value in REPEAT_CELLS.items():
        cell = sheet[coordinate]
        cell.value = value
        if coordinate == "E19":  # was the text column "T/E", now money
            cell.number_format = currency_format

    sheet.delete_rows(FIRST_DATA_ROW, LAST_TOTALS_ROW - FIRST_DATA_ROW + 1)
    sheet.insert_rows(FIRST_DATA_ROW, len(TOTALS))
    for offset, (label, placeholder, _) in enumerate(TOTALS):
        row = FIRST_DATA_ROW + offset
        for column, value in ((6, label), (7, placeholder)):
            cell = sheet.cell(row, column)
            _copy_style(totals_styles[offset * 2 + (column - 6)], cell)
            cell.value = value

    for index in list(sheet.row_dimensions):
        sheet.row_dimensions[index].height = None
    for index, height in heights.items():
        if index < FIRST_DATA_ROW:
            sheet.row_dimensions[index].height = height
        elif index > LAST_TOTALS_ROW:
            sheet.row_dimensions[index - SHIFT].height = height
    for offset, (_, _, source_row) in enumerate(TOTALS):
        height = heights.get(source_row)
        if height is not None:
            sheet.row_dimensions[FIRST_DATA_ROW + offset].height = height

    for reference in MERGES_ABOVE:
        sheet.merge_cells(reference)
    sheet.merge_cells(TOTALS_BLOCK)
    for reference in MERGES_BELOW:
        sheet.merge_cells(_shift_range(reference, -SHIFT))

    if with_letterhead:
        _anchor_letterhead(source, sheet)

    workbook.save(output)
    workbook.close()


def _anchor_letterhead(source: Path, sheet: Worksheet) -> None:
    from openpyxl.drawing.image import Image

    with zipfile.ZipFile(source) as archive:
        payload = archive.read("xl/media/image1.jpeg")
    image = Image(io.BytesIO(payload))
    image.width = LETTERHEAD_WIDTH_PT * 96 / 72
    image.height = LETTERHEAD_HEIGHT_PT * 96 / 72
    sheet.add_image(image, "A1")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--with-letterhead", action="store_true")
    arguments = parser.parse_args()
    build(arguments.source, arguments.output, with_letterhead=arguments.with_letterhead)
    print(f"Plantilla escrita en {arguments.output}")


if __name__ == "__main__":
    main()
